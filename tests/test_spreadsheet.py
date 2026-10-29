import copy
import io
from pathlib import Path
import tempfile
from threading import Event
from types import SimpleNamespace
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from chickadee.spreadsheet import COLUMNS, export_rows, save_results, xlsx_bytes

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def sample():
    result = {'meta': {'tool_version': 'test'}, 'findings': [
        {'id': 'f1', 'url': 'https://example.test/post/1', 'location': 'article',
         'evidence_text': '=HYPERLINK("https://example.test","광고")\n한글 원문', 'technique': 'HOMOGLYPH', 'is_violation': False}]}
    report = {'evidence_details': {'f1': {'source_kind': 'dom_text', 'evidence_sources': [
        {'source_kind': 'image_ocr', 'image_id': 'im1'}]}},
        'ocr_images': [{'id': 'im1', 'text': '한국어 이미지 문구'}],
        'ad_candidates': [{'id': 'c1', 'url': 'https://example.test/post/1#comment', 'location': 'article .comment',
            'evidence_text': '댓글 광고', 'category': 'GAMBLING'}]}
    reviews = {'items': {'f1': {'status': 'confirmed', 'note': '확인한 근거', 'updated_at': '2026-10-29T12:00:00+09:00'}}}
    return result, report, reviews


class SpreadsheetTests(unittest.TestCase):
    def test_korean_ocr_review_literal_text_and_original_preserved(self):
        args = sample()
        original = copy.deepcopy(args)
        payload = xlsx_bytes(*args)
        self.assertEqual(args, original)
        with ZipFile(io.BytesIO(payload)) as archive:
            self.assertIsNone(archive.testzip())
            for name in archive.namelist():
                ET.fromstring(archive.read(name))
            sheet = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(sheet.find('m:autoFilter', NS).get('ref'), 'A1:L3')
            self.assertEqual(sheet.find('m:sheetViews/m:sheetView/m:pane', NS).get('state'), 'frozen')
            self.assertEqual(sheet.findall('.//m:f', NS), [])
            self.assertEqual(sheet.findall('.//m:hyperlink', NS), [])
            self.assertEqual(sheet.find('.//m:c[@r="E2"]/m:is/m:t', NS).text, args[0]['findings'][0]['evidence_text'])
        rows = list(export_rows(*args))
        self.assertEqual(rows[0][1][6:10], ['이미지 OCR', '한국어 이미지 문구', '검토 확인', '확인한 근거'])

    def test_independent_excel_reader(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('optional independent XLSX reader is unavailable')
        book = load_workbook(io.BytesIO(xlsx_bytes(*sample())))
        sheet = book['탐지 내역']
        self.assertEqual(tuple(c.value for c in sheet[1]), COLUMNS)
        self.assertEqual(sheet['E2'].data_type, 's')
        self.assertEqual(sheet['H2'].value, '한국어 이미지 문구')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.max_row, 3)
        book.close()

    def test_photo_classification_does_not_claim_ocr(self):
        result, report, reviews = sample()
        entry = report['ad_candidates'][0]
        entry.update(source_kind='image_visual', image_id='im1')
        row = list(export_rows(result, report, reviews))[1][1]
        self.assertEqual(row[6], '이미지 분석')
        self.assertEqual(row[7], '')

    def test_selected_csv_xlsx_cancelled_write_and_bad_ids(self):
        args = sample()
        with tempfile.TemporaryDirectory() as folder:
            for extension in ('xlsx', 'csv'):
                path = Path(folder)/('선택.' + extension)
                value = save_results(path, *args, ids=['c1', 'c1'])
                self.assertEqual(value['count'], 1)
                if extension == 'xlsx':
                    with ZipFile(path) as archive:
                        self.assertIn('댓글 광고', archive.read('xl/worksheets/sheet1.xml').decode())
                        self.assertNotIn('한글 원문', archive.read('xl/worksheets/sheet1.xml').decode())
                else:
                    self.assertNotIn('한글 원문', path.read_text(encoding='utf-8-sig'))
                before = path.read_bytes()
                for invalid in ([], ['unknown'], 'f1', [1]):
                    with self.assertRaises(ValueError):
                        save_results(path, *args, ids=invalid)
                    self.assertEqual(path.read_bytes(), before)
            with self.assertRaises(ValueError):
                save_results(Path(folder)/'wrong.html', *args)

    def test_xml_control_characters_and_excel_cell_limit(self):
        result, report, reviews = sample()
        result['findings'][0]['evidence_text'] = '한글\x00\x08\x1f<&>😀'
        with ZipFile(io.BytesIO(xlsx_bytes(result, report, reviews))) as archive:
            ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        result['findings'][0]['evidence_text'] = '가' * 32768
        with self.assertRaisesRegex(ValueError, '32767'):
            xlsx_bytes(result, report)
        empty = xlsx_bytes({'meta': {}, 'findings': []}, {})
        with ZipFile(io.BytesIO(empty)) as archive:
            sheet = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(len(sheet.findall('m:sheetData/m:row', NS)), 1)

    def test_bridge_rejects_stale_run_and_busy_scan(self):
        from chickadee.bridge import DesktopBridge
        bridge = DesktopBridge.__new__(DesktopBridge)
        bridge.busy = Event()
        result, report, reviews = sample()
        bridge.outcome = SimpleNamespace(result=result, report=report)
        bridge.reviews = SimpleNamespace(data=reviews)
        bridge.run_id = 'run1'
        with tempfile.TemporaryDirectory() as folder:
            bridge.data_dir = Path(folder)
            path = Path(folder)/'결과.xlsx'
            with self.assertRaises(ValueError):
                bridge.command('export', {'path': str(path), 'run_id': 'run2'})
            self.assertFalse(path.exists())
            self.assertEqual(bridge.command('export', {'path': str(path), 'run_id': 'run1', 'ids': ['f1']})['count'], 1)
            bridge.busy.set()
            with self.assertRaises(ValueError):
                bridge.command('export', {'path': str(path), 'run_id': 'run1'})


if __name__ == '__main__':
    unittest.main()
